"""Compare the two entries of every survey typed by both people → data/working/differences.xlsx.

Run from the project root (after `python -m recon.questions`):  python -m recon.compare
"""

import json
import unicodedata

import pandas as pd
from openpyxl.styles import Font, PatternFill

from recon.config import load_config
from recon.loader import load

# Typed-text differences that don't change the answer.
_QUOTES = str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"', "​": None, "﻿": None})


def same_text(value: str) -> str:
    """Typed text with spacing, capitals, quote style and invisible characters evened out.
    Punctuation and spelling are kept, so those differences are still listed."""
    return " ".join(unicodedata.normalize("NFKC", value).translate(_QUOTES).split()).casefold()


def _label(value: str, column: dict, question: dict) -> str:
    """An answer as the team reads it, e.g. "2 = 1-5 years"."""
    if value == "":
        return "(blank)"
    if column["is_text"]:
        return value
    if question["kind"] == "check_all":
        return "ticked" if value == "1" else value
    labels = question.get("scale" if question["kind"] == "grid" else "choices", {})
    return f"{value} = {labels[value]}" if value in labels else value


def compare(a: pd.DataFrame, b: pd.DataFrame, qmap: dict, names=("A", "B")) -> pd.DataFrame:
    """One row per disagreement between the two entries of each survey typed by both people."""
    na, nb = names
    both = a.merge(b, on="survey_code", suffixes=("_A", "_B"))
    rows = []
    for column in (c for c in qmap["columns"] if c["role"] == "COMPARE"):
        q = qmap["questions"][column["qid"]]
        x, y = both[column["column"] + "_A"], both[column["column"] + "_B"]
        differ = x.map(same_text) != y.map(same_text) if column["is_text"] else x != y
        option = {**q.get("choices", {}), **q.get("rows", {})}.get(column["choice"], "")
        question = f"{q['tag']}: {q['text']}" + (f" | {option}" if option else "")
        for _, r in both[differ].iterrows():
            rows.append({
                "survey_code": r["survey_code"], "column": column["column"], "question": question,
                f"{na} answer": _label(r[column["column"] + "_A"], column, q),
                f"{nb} answer": _label(r[column["column"] + "_B"], column, q),
                f"{na} ResponseId": r["ResponseId_A"], f"{nb} ResponseId": r["ResponseId_B"],
            })
    return pd.DataFrame(rows, columns=[
        "survey_code", "column", "question", f"{na} answer", f"{nb} answer",
        f"{na} ResponseId", f"{nb} ResponseId"]).sort_values("survey_code", kind="stable")


# Set-aside reasons from the loader, in plain words.
_REASONS = {"REPEATED_IN_EXPORT": "typed more than once", "NOT_ON_PAPER_LOG": "not in paper log",
            "REPEATED_ON_PAPER_LOG": "twice in paper log",
            "CODE_FORMAT_NOT_CHECKABLE_AGAINST_PAPER_LOG": "unusual code"}


def _plain(reason: str) -> str:
    """'REPEATED_IN_EXPORT(2x); NOT_ON_PAPER_LOG' -> 'set aside (typed more than once, not in paper log)'"""
    words = [_REASONS.get(r.split("(")[0], r) for r in reason.split("; ")]
    return f"set aside ({', '.join(words)})"


def survey_table(a, b, excluded, differences, names=("A", "B")) -> pd.DataFrame:
    """One row per survey code: whether each person's entry was used, and the comparison result."""
    na, nb = names
    state = {}
    for name, used in ((na, a), (nb, b)):
        for code in used["survey_code"]:
            state.setdefault(code, {})[name] = "typed"
        for code, reason in zip(*(excluded.loc[excluded["entrant"] == name, c] for c in ("survey_code", "reason"))):
            state.setdefault(code, {}).setdefault(name, _plain(reason))
    counts = differences.groupby("survey_code").size()
    where = differences.groupby("survey_code")["column"].agg(", ".join)
    rows = []
    for code, s in sorted(state.items()):
        x, y = s.get(na, "not typed"), s.get(nb, "not typed")
        if x == y == "typed":
            n = int(counts.get(code, 0))
            status = "Compared: all answers match" if n == 0 else f"Compared: {n} difference{'s' * (n > 1)}"
        else:
            status = f"Not compared: {na} {x}, {nb} {y}"
        rows.append({"survey_code": code, "status": status, f"{na} entry": x, f"{nb} entry": y,
                     "differences": int(counts.get(code, 0)), "columns that differ": where.get(code, "")})
    return pd.DataFrame(rows)


def _write(xl, sheet: str, table: pd.DataFrame):
    """A sheet with a filtered, frozen header row and readable column widths."""
    table.to_excel(xl, sheet_name=sheet, index=False)
    ws = xl.sheets[sheet]
    ws.auto_filter.ref, ws.freeze_panes = ws.dimensions, "A2"
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    for cells in ws.columns:
        width = max(len(str(c.value or "")) for c in cells)
        ws.column_dimensions[cells[0].column_letter].width = min(width + 2, 60)


def main():
    config = load_config()
    qmap = json.loads((config.path.parent / "questions.json").read_text(encoding="utf-8"))
    na, nb = (next(i for i, e in config.entrants.items() if e.dataset == d) for d in ("A", "B"))
    result = load(config)
    a, b = result.datasets["A"], result.datasets["B"]
    differences = compare(a, b, qmap, (na, nb))
    surveys = survey_table(a, b, result.excluded, differences, (na, nb))
    compared = surveys["status"].str.startswith("Compared")
    by_question = (differences.groupby(["column", "question"], sort=False).size()
                   .rename("differences").reset_index().sort_values("differences", ascending=False))
    summary = pd.DataFrame([
        ("Surveys compared (typed by both)", compared.sum()),
        ("  all answers match", (surveys["differences"] == 0)[compared].sum()),
        ("  with at least one difference", (surveys["differences"] > 0).sum()),
        ("Differences to check (rows on the Differences sheet)", len(differences)),
        (f"Surveys not compared: only {na} typed it (or the other entry is set aside)", ((surveys[f"{na} entry"] == "typed") & ~compared).sum()),
        (f"Surveys not compared: only {nb} typed it (or the other entry is set aside)", ((surveys[f"{nb} entry"] == "typed") & ~compared).sum()),
        ("Surveys not compared: both entries set aside or missing",
         ((surveys[f"{na} entry"] != "typed") & (surveys[f"{nb} entry"] != "typed")).sum()),
        ("Entries with no initials (not in this report)", (result.excluded["entrant"] == "").sum()),
        ("Same: extra spaces, capital letters, line breaks, quote style", ""),
        ("Different: any other change, including punctuation, spelling, and answer vs blank", ""),
    ], columns=["What", "Count"])

    out = config.working_dir / "differences.xlsx"
    with pd.ExcelWriter(out) as xl:
        for sheet, table in (("Summary", summary), ("Surveys", surveys),
                             ("Differences", differences), ("By question", by_question)):
            _write(xl, sheet, table)
    print(f"Compared {compared.sum()} surveys: {len(differences)} differences in "
          f"{(surveys['differences'] > 0).sum()} surveys. Wrote {out}")


if __name__ == "__main__":
    main()
